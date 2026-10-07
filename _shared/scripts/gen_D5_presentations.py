#!/usr/bin/env python3
"""
Д+5: Сборка материалов АС Фонд-М + презентации (PPTX) для 6 тем.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

REPO = "/home/user/The-competition-is-for-5-applications-all-6-are-accepted."

TOPICS = [
    {
        "dir": "01_termokist-6", "short": "ТермоКисть-6",
        "full": "Портативная система управляемого нагрева с PCM и цифровым управлением",
        "leader": "Папуша С.К.",
        "budget": "1 040 000",
        "trl_start": "3", "trl_end": "6",
        "market": "Системы локального нагрева: RL-2600 (2795₽), Alpenheat (19500–25000₽)",
        "target": "3500 ₽",
        "key_results": [
            "RC-модель теплового баланса (C·dT/dt)",
            "PCM-аккумулятор (T_пл = 42 °C, NiCr нагреватель)",
            "Корпус TPU 95A + аэрогелевая изоляция",
            "BLE 5.0, мобильное приложение",
            "MTBF: 120 482 ч",
        ],
        "advantage": "Стоимость ≤3500 ₽ (vs 19 500 ₽ Alpenheat); цифровое управление; PCM не требует замены",
    },
    {
        "dir": "02_vetpcr-chemodan", "short": "ВетПЦР-Чемодан",
        "full": "Портативная ПЦР-лаборатория для полевой диагностики ветеринарных инфекций",
        "leader": "Инюкина Т.А.",
        "budget": "1 250 000",
        "trl_start": "3", "trl_end": "6",
        "market": "Портативные ПЦР: Biomeme Two3 (~$3500), miniPCR (~$650)",
        "target": "45 000 ₽",
        "key_results": [
            "Термоблок с Пельтье (55–95 °C, ≥4 °C/с)",
            "4-канальный детектор (FAM/HEX/ROX/Cy5)",
            "Охлаждение реагентов 4±2 °C (PCM)",
            "GPS-привязка + BLE",
            "MTBF: 65 359 ч",
        ],
        "advantage": "Время от пробы до результата ≤90 мин (vs 3–5 суток); 4 канала; GPS-привязка",
    },
    {
        "dir": "03_parazit-kontrol", "short": "Паразит-Контроль",
        "full": "Система компьютерного зрения для подсчёта личинок паразитов в рыбной продукции",
        "leader": "Кремянский В.Ф.",
        "budget": "950 000",
        "trl_start": "2", "trl_end": "5",
        "market": "Аналогов для подсчёта личинок не обнаружено",
        "target": "120 000 ₽",
        "key_results": [
            "YOLOv8n: mAP@0.5 ≥ 0.85",
            "Синтетический датасет 200+ изображений",
            "Камера IMX477 + макрообъектив 25 мм",
            "Точность ≥92%, время ≤2 мин/образец",
            "MTBF: 83 333 ч",
        ],
        "advantage": "Пропускная способность ≥25 образцов/час (vs 3–5); воспроизводимость; протокол экспертизы",
    },
    {
        "dir": "04_optikozor-2-0", "short": "ОптикоЗор-2,0",
        "full": "Высокоскоростная система машинного зрения для инспекции полупроводниковых пластин",
        "leader": "Богус А.Э.",
        "budget": "1 500 000",
        "trl_start": "3", "trl_end": "6",
        "market": "KLA-Tencor Surfscan (>$5M), Lasertec MA-T (>$5M)",
        "target": "45 000 000 ₽ (~$500K)",
        "key_results": [
            "4K CMOS-линейка (1 мкм/пиксель)",
            "Гиперспектральный осветитель (400–900 нм)",
            "FPGA Zynq UltraScale+",
            "PCA + автоэнкодеры для детекции",
            "MTBF: ≥8 000 ч",
        ],
        "advantage": "Стоимость ≤$500K (vs >$5M); гиперспектральный анализ; ≥50 пластин/час",
    },
    {
        "dir": "05_cehgid-5", "short": "ЦехГид-5",
        "full": "Навигационная система автономного робота для мониторинга микроклимата",
        "leader": "Класнер Г.Г.",
        "budget": "1 100 000",
        "trl_start": "3", "trl_end": "6",
        "market": "Fancom F48 (стационарный), Big Dutchman BigFarmNet",
        "target": "150 000 ₽",
        "key_results": [
            "RPLIDAR A3 (360°, 16 м, 15 Гц)",
            "SLAM Cartographer на Raspberry Pi 4B",
            "Блок сенсоров: SHT45, SGP41, SCD41, TSL2591",
            "LoRa SX1276 (868 МГц, до 5 км)",
            "MTBF: 117 647 ч",
        ],
        "advantage": "Пространственное разрешение ≤2 м (vs 3–5 точек на цех); ≤45 мин обход; снижение стресса птицы",
    },
    {
        "dir": "06_smena-prognoz-40", "short": "Смена-Прогноз-40",
        "full": "Носимая система мониторинга с прогнозом теплового стресса за 40 минут",
        "leader": "Самурганов Е.Е.",
        "budget": "980 000",
        "trl_start": "3", "trl_end": "6",
        "market": "Kenzen ($300–500), SlateSafety ($200–500)",
        "target": "3000 ₽",
        "key_results": [
            "nRF52832 + AFE4404 + MLX90614 + AD5940",
            "CatBoost: прогноз за 40 мин, точность ≥85%",
            "PPG + КГР + температура кожи + акселерометр",
            "Корпус силикон, IP67",
            "MTBF: ≥4 000 ч",
        ],
        "advantage": "Стоимость ≤3000 ₽ (vs $200–500); прогноз за 40 мин; интеграция с диспетчерской",
    },
]


def create_pptx(t):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Slide 1: Title
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11), Inches(2))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = t["short"]
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    p2 = tf.add_paragraph()
    p2.text = t["full"]
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(0x37, 0x41, 0x51)

    p3 = tf.add_paragraph()
    p3.text = f"\nФСИ «Старт-Пром-1», очередь 4 | Руководитель: {t['leader']}"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    # Slide 2: Problem & Goal
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(1))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Проблема и цель"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    txBox2 = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(5.5), Inches(5))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "Рынок"
    p.font.size = Pt(18); p.font.bold = True
    p2 = tf2.add_paragraph()
    p2.text = t["market"]
    p2.font.size = Pt(14)
    p3 = tf2.add_paragraph()
    p3.text = f"\nЦелевая стоимость: {t['target']}"
    p3.font.size = Pt(14); p3.font.bold = True

    txBox3 = slide.shapes.add_textbox(Inches(6.5), Inches(1.5), Inches(6), Inches(5))
    tf3 = txBox3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "Преимущества"
    p.font.size = Pt(18); p.font.bold = True
    p2 = tf3.add_paragraph()
    p2.text = t["advantage"]
    p2.font.size = Pt(14)

    # Slide 3: Key Results
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(1))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Ключевые результаты"
    p.font.size = Pt(32); p.font.bold = True
    p.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    txBox2 = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12), Inches(5))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    for i, kr in enumerate(t["key_results"]):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.text = f"✓ {kr}"
        p.font.size = Pt(18)
        p.space_after = Pt(10)

    # Slide 4: Budget & Timeline
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(1))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Бюджет и этапы"
    p.font.size = Pt(32); p.font.bold = True
    p.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    txBox2 = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12), Inches(5))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = f"Общий бюджет: {t['budget']} руб."
    p.font.size = Pt(24); p.font.bold = True

    p2 = tf2.add_paragraph()
    p2.text = f"Срок: 18 месяцев | TRL: {t['trl_start']} → {t['trl_end']}"
    p2.font.size = Pt(18)

    p3 = tf2.add_paragraph()
    p3.text = "\nЭтап 1 (мес. 1–8): Разработка прототипа"
    p3.font.size = Pt(16); p3.font.bold = True
    p4 = tf2.add_paragraph()
    p4.text = "Этап 2 (мес. 9–14): Опытная партия и испытания"
    p4.font.size = Pt(16)
    p5 = tf2.add_paragraph()
    p5.text = "Этап 3 (мес. 15–18): Документация и передача результатов"
    p5.font.size = Pt(16)

    p6 = tf2.add_paragraph()
    p6.text = "\nПатентование РИД финансируется из собственных средств заявителя и не включено в бюджет гранта."
    p6.font.size = Pt(12); p6.font.italic = True
    p6.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    out_path = os.path.join(REPO, t["dir"], "05_zayavka_AS", f"prezentaciya_{t['dir'].split('_')[0]}.pptx")
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs.save(out_path)
    return out_path


# ── AS Фонд-М сводка ──
def create_fond_summary():
    out = os.path.join(REPO, "_shared/AS_FOND_M_SVODKA.md")
    with open(out, 'w', encoding='utf-8') as f:
        f.write("# Сводка материалов АС Фонд-М (6 заявок)\n\n")
        f.write("**Конкурс:** ФСИ «Старт-Пром-1», очередь 4\n")
        f.write("**Дедлайн:** 10:00 мск 12.10.2026\n\n")
        f.write("| № | Тема | Руководитель | Бюджет | TRL | Статус |\n")
        f.write("|---|------|-------------|--------|-----|--------|\n")
        for t in TOPICS:
            f.write(f"| {t['dir'][:2]} | {t['short']} | {t['leader']} | {t['budget']} ₽ | {t['trl_start']}→{t['trl_end']} | ✅ |\n")
        f.write(f"\n**Итого бюджет:** {sum(int(t['budget'].replace(' ','')) for t in TOPICS):,} ₽\n\n")
        f.write("## Комплект документов на каждую заявку\n\n")
        f.write("1. Паспорт темы (README.md)\n")
        f.write("2. Расчёты (model.py + input_data.csv + XLSX + графики)\n")
        f.write("3. Чертежи (DXF + PDF + STL)\n")
        f.write("4. BOM + спецификация\n")
        f.write("5. Электрическая схема\n")
        f.write("6. ТЗ (техническое задание)\n")
        f.write("7. КП (коммерческое предложение / бюджет)\n")
        f.write("8. Заявка на ПМ (описание + формула + реферат + чертежи)\n")
        f.write("9. Заявка на регистрацию ПО\n")
        f.write("10. Протоколы (стендовых + натурных испытаний, шаблоны)\n")
        f.write("11. Акт приёмки этапа\n")
        f.write("12. Отчёт по этапам\n")
        f.write("13. ЕГИСУ НИОКТР (форма)\n")
        f.write("14. План реализации\n")
        f.write("15. Публикация (рукопись)\n")
        f.write("16. Письма партнёрам\n")
        f.write("17. Документы грантового договора (Прил. 7–10, 2 соглашения)\n")
        f.write("18. Таблица СМИ-источников\n")
        f.write("19. Презентация (PPTX)\n")
    print(f"✓ AS Фонд-М сводка: {out}")


create_fond_summary()

for t in TOPICS:
    path = create_pptx(t)
    print(f"  {t['short']}: презентация ✓ ({os.path.basename(path)})")

print("\n✓ Д+5: презентации и сводка АС Фонд-М.")